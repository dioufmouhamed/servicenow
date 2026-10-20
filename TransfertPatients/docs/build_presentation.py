#!/usr/bin/env python3
"""Génère docs/Transfert_Patients_presentation.pptx (python-pptx)."""
import os
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Transfert_Patients_presentation.pptx")

TEAL, TEAL_D, MINT, AMBER = "0B5D57", "083F3B", "E4F1EE", "E39B2D"
INK, GREY, WHITE, LINE = "1E2B2E", "5A6B6F", "FFFFFF", "C9D8D5"
HEAD, BODY = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = rgb(color)


def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None, spacing=6):
    """runs: str | list de paragraphes ; un paragraphe = str ou (str, {options})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        t, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = o.get("align", align)
        para.space_after = Pt(o.get("spacing", spacing))
        r = para.add_run()
        r.text = t
        r.font.size = Pt(o.get("size", size))
        r.font.bold = o.get("bold", bold)
        r.font.italic = o.get("italic", False)
        r.font.name = o.get("font", font)
        r.font.color.rgb = rgb(o.get("color", color))
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def badge(slide, x, y, label, d=0.5, fill=TEAL, color=WHITE, size=16):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(size), True, BODY
    r.font.color.rgb = rgb(color)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def arrow(slide, x1, y1, x2, y2, color=GREY):
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
    ln.append(tail)
    return c


def content_slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    if kicker:
        text(s, 0.6, 0.38, 12, 0.3, kicker.upper(), size=12, color=TEAL, bold=True, name="Kicker")
    text(s, 0.6, 0.66, 12.1, 0.9, title, size=32, font=HEAD, bold=True, name="Titre")
    text(s, 12.2, 7.05, 0.6, 0.3, str(len(prs.slides)), size=10, color=GREY, align=PP_ALIGN.RIGHT, name="Numéro")
    return s


# 1 — Titre ---------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, TEAL_D)
badge(s, 0.8, 1.2, "+", d=0.9, fill=AMBER, color=TEAL_D, size=36)
text(s, 0.8, 2.45, 11.5, 1.3, "Transfert Patients", size=54, font=HEAD, bold=True, color=WHITE)
text(s, 0.8, 3.6, 11.5, 1.0, "Transferts inter-établissements de patients sur ServiceNow, "
     "avec Field Service Management — région de Dakar", size=22, color="CFE6E2")
text(s, 0.8, 5.9, 11.5, 0.8, ["Application scopée x_2218203_ptrans · PDI dev408543",
                              "Octobre 2026"], size=14, color="9CC5BF")
s.notes_slide.notes_text_frame.text = (
    "Présentation de la solution Transfert Patients : objectif, processus, architecture ServiceNow et FSM, "
    "sécurité, démonstration et feuille de route.")

# 2 — Le besoin -------------------------------------------------------------
s = content_slide("Un établissement publie, les autres décident, FSM transporte", "Le besoin et la réponse")
cards = [
    ("1", "Publier", "Le régulateur décrit le besoin du patient. La demande part seulement aux établissements "
                     "qui ont la spécialité, un lit libre et les équipements requis."),
    ("2", "Décider librement", "Chaque établissement accepte (en proposant un service) ou refuse "
                               "(avec un motif). Le régulateur choisit parmi les acceptations."),
    ("3", "Transporter", "Le choix réserve un lit et crée un Work Order Field Service : qualification, "
                         "dispatch par zone, exécution par l'équipage."),
]
for i, (n, t, d) in enumerate(cards):
    x = 0.6 + i * 4.1
    box(s, x, 2.0, 3.8, 4.3, fill=MINT, name=f"Carte {t}")
    badge(s, x + 0.35, 2.35, n)
    text(s, x + 0.35, 3.1, 3.1, 0.5, t, size=22, font=HEAD, bold=True, color=TEAL_D)
    text(s, x + 0.35, 3.75, 3.1, 2.4, d, size=15)
s.notes_slide.notes_text_frame.text = "Trois temps : publication ciblée, décision libre de chaque établissement, transport OOTB FSM."

# 3 — Chiffres clés ------------------------------------------------------------
s = content_slide("Une application légère posée sur des briques standard", "En chiffres")
stats = [("5", "tables spécifiques", "le reste est OOTB : localisations, assets, FSM, groupes"),
         ("9", "Business Rules", "+ 1 Script Include et 5 boutons"),
         ("7", "hôpitaux de Dakar", "dans les données de démonstration"),
         ("3", "zones de dispatch", "Plateau-Médina, Fann-Grand Yoff, Pikine-Guédiawaye")]
for i, (v, l, d) in enumerate(stats):
    x = 0.6 + i * 3.1
    text(s, x, 2.2, 2.9, 1.3, v, size=66, font=HEAD, bold=True, color=TEAL)
    text(s, x, 3.6, 2.9, 0.5, l, size=18, bold=True)
    text(s, x, 4.15, 2.7, 1.2, d, size=14, color=GREY)
box(s, 0.6, 5.7, 12.1, 0.9, fill=MINT)
text(s, 0.9, 5.7, 11.6, 0.9, "Principe : FSM reste intact. Les Work Orders avancent uniquement par les "
     "State Flows OOTB, selon la configuration de l'instance.", size=16, color=TEAL_D, anchor=MSO_ANCHOR.MIDDLE)

# 4 — Processus -------------------------------------------------------------------
s = content_slide("Six étapes de la création à l'arrivée du patient", "Processus métier")
steps = [("Brouillon", "Le régulateur saisit la demande"), ("Publiée", "Réponses créées pour les éligibles"),
         ("Réponses", "Accepter avec un service ou refuser avec un motif"),
         ("Destination retenue", "Lit réservé, autres notifiés"), ("Transport planifié", "Work Order FSM créé"),
         ("Arrivé", "WO clôturé, demande clôturée")]
w, gap, y = 1.85, 0.2, 2.3
for i, (t, d) in enumerate(steps):
    x = 0.6 + i * (w + gap)
    last = i == len(steps) - 1
    b = box(s, x, y, w, 1.1, fill=TEAL if last else MINT, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON,
            name=f"Étape {i + 1}")
    text(s, x + 0.3, y, w - 0.5, 1.1, t, size=14, bold=True, color=WHITE if last else TEAL_D,
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    text(s, x + 0.1, y + 1.3, w - 0.2, 1.2, d, size=13, color=GREY, align=PP_ALIGN.CENTER)
box(s, 0.6, 4.9, 5.9, 1.6, fill="FCF1DF")
text(s, 0.9, 5.05, 5.4, 1.4, [("Expiration", {"bold": True, "size": 16, "color": "8A5A12"}),
                              ("Sans acceptation avant la date limite (30 min, 2 h ou 24 h selon l'urgence), "
                               "le job planifié passe la demande en Expirée.", {"size": 14})])
box(s, 6.8, 4.9, 5.9, 1.6, fill=MINT)
text(s, 7.1, 5.05, 5.4, 1.4, [("Annulation", {"bold": True, "size": 16, "color": TEAL_D}),
                              ("Possible en Brouillon, Publiée ou Destination retenue. "
                               "Les réponses ouvertes sont clôturées.", {"size": 14})])

# 5 — Éligibilité ---------------------------------------------------------------------
s = content_slide("Seuls les établissements capables de soigner reçoivent la demande", "Règle d'éligibilité")
crit = [("Actif et ouvert", "L'établissement est actif et accepte les transferts"),
        ("Pas l'origine", "L'établissement d'origine est exclu"),
        ("Spécialité et lit", "Un service actif de la spécialité requise a au moins 1 lit libre"),
        ("Équipements", "Chaque équipement requis existe sur place (asset en service ou en stock)")]
for i, (t, d) in enumerate(crit):
    y = 1.95 + i * 1.15
    badge(s, 0.6, y, str(i + 1))
    text(s, 1.35, y - 0.02, 5.2, 0.4, t, size=18, bold=True)
    text(s, 1.35, y + 0.4, 5.2, 0.6, d, size=14, color=GREY)
box(s, 7.2, 1.95, 5.5, 4.55, fill=MINT, name="Exemple")
text(s, 7.55, 2.2, 4.9, 4.2, [
    ("Exemple : TRF0001013", {"bold": True, "size": 18, "color": TEAL_D, "font": HEAD}),
    ("Origine : Hôpital Dalal Jamm", {"size": 15}),
    ("Besoin : pédiatrie, soins continus, moniteur", {"size": 15, "spacing": 14}),
    ("Résultat : publiée à 1 établissement", {"bold": True, "size": 15}),
    ("Hôpital d'Enfants Albert Royer, service Pédiatrie (6 lits libres).", {"size": 15, "spacing": 14}),
    ("Le service proposé est celui qui a le plus de lits libres.", {"size": 13, "color": GREY, "italic": True}),
])

# 6 — Rôles -----------------------------------------------------------------------------
s = content_slide("Six profils, trois rôles applicatifs et trois rôles FSM", "Rôles et personas")
roles = [("Régulateur", "Crée, publie, retient la destination", "reg.demo"),
         ("Coordinateur", "Gère les lits, accepte ou refuse pour son établissement", "coord.fann, coord.principal…"),
         ("Admin fonctionnel", "Gère établissements, services, spécialités", "rôle x_2218203_ptrans.admin"),
         ("Qualificateur FSM", "Qualifie les ordres de transport", "wm_qualifier"),
         ("Dispatcher FSM", "Affecte les tâches aux équipages de sa zone", "disp.plateau, disp.fann…"),
         ("Ambulancier", "Exécute les tâches, y compris sur mobile", "wm_agent")]
for i, (t, d, u) in enumerate(roles):
    col, row = i % 3, i // 3
    x, y = 0.6 + col * 4.1, 1.95 + row * 2.35
    app = i < 3
    box(s, x, y, 3.8, 2.1, fill=MINT if app else WHITE, line=None if app else LINE, name=f"Rôle {t}")
    text(s, x + 0.3, y + 0.25, 3.2, 0.45, t, size=18, bold=True, color=TEAL_D)
    text(s, x + 0.3, y + 0.75, 3.2, 0.8, d, size=14)
    text(s, x + 0.3, y + 1.6, 3.2, 0.35, u, size=12, color=GREY)
text(s, 0.6, 6.75, 12, 0.3, "Fond teinté : rôles de l'application · fond blanc : rôles Field Service Management standard",
     size=11, color=GREY)

# 7 — Modèle de données --------------------------------------------------------------------
s = content_slide("Cinq tables spécifiques, reliées aux tables standard", "Modèle de données")
def ent(x, y, w, title, sub, ootb=False):
    box(s, x, y, w, 0.95, fill=WHITE if ootb else MINT, line=LINE if ootb else None, name=f"Table {title}")
    text(s, x + 0.2, y + 0.12, w - 0.4, 0.4, title, size=15, bold=True, color=GREY if ootb else TEAL_D)
    text(s, x + 0.2, y + 0.5, w - 0.4, 0.4, sub, size=11, color=GREY)
ent(0.6, 2.0, 3.0, "Spécialité", "référentiel")
ent(0.6, 3.55, 3.0, "Service", "lits totaux / disponibles")
ent(0.6, 5.1, 3.0, "Établissement", "capacité calculée")
ent(5.15, 2.0, 3.0, "Demande (TRF)", "étend task · 8 états")
ent(5.15, 3.55, 3.0, "Réponse (REP)", "décision par établissement")
ent(9.7, 2.0, 3.0, "wm_order / wm_task", "ordre de transport FSM", True)
ent(9.7, 3.55, 3.0, "cmdb_model / alm_asset", "équipements médicaux", True)
ent(9.7, 5.1, 3.0, "cmn_location", "adresse, GPS, zone", True)
arrow(s, 2.1, 3.55, 2.1, 2.95)
arrow(s, 2.1, 4.5, 2.1, 5.1)
arrow(s, 6.65, 3.55, 6.65, 2.95)
arrow(s, 5.15, 4.0, 3.6, 5.45)
arrow(s, 8.15, 2.47, 9.7, 2.47)
arrow(s, 3.6, 5.57, 9.7, 5.57)
text(s, 5.15, 5.1, 4.2, 1.2, "Les flèches vont de la table qui porte la référence vers la table référencée. "
     "Les équipements d'un établissement sont les assets situés à sa localisation.", size=12, color=GREY)

# 8 — Intégration FSM ----------------------------------------------------------------------
s = content_slide("Le transport est un Work Order FSM standard à trois tâches", "Intégration Field Service")
tasks = [("Préparation du patient", "à l'origine · identité, dossier, traitements"),
         ("Transport", "départ de l'origine · détail médical"),
         ("Remise du patient", "à destination · transmission et signature")]
box(s, 0.6, 2.0, 5.4, 0.9, fill=TEAL, name="Work Order")
text(s, 0.9, 2.0, 5.0, 0.9, "Work Order · priorité selon l'urgence", size=16, bold=True, color=WHITE,
     anchor=MSO_ANCHOR.MIDDLE)
for i, (t, d) in enumerate(tasks):
    y = 3.2 + i * 1.1
    box(s, 1.2, y, 4.8, 0.9, fill=MINT, name=f"Tâche {i + 1}")
    text(s, 1.45, y + 0.1, 4.4, 0.4, f"Tâche {i + 1} · {t}", size=14, bold=True, color=TEAL_D)
    text(s, 1.45, y + 0.48, 4.4, 0.4, d, size=12, color=GREY)
text(s, 6.7, 2.0, 6.0, 4.6, [
    ("Ce que fait l'application", {"bold": True, "size": 18, "color": TEAL_D, "font": HEAD}),
    ("Elle crée l'ordre et ses tâches, puis applique le premier State Flow OOTB autorisé par la "
     "configuration FSM.", {"size": 15, "spacing": 12}),
    ("Ce que fait FSM", {"bold": True, "size": 18, "color": TEAL_D, "font": HEAD}),
    ("Qualification active sur l'instance : l'ordre passe en Awaiting Qualification. FSM renseigne seul "
     "le groupe de qualification et le groupe de dispatch de chaque tâche.", {"size": 15, "spacing": 12}),
    ("Retour vers la demande", {"bold": True, "size": 18, "color": TEAL_D, "font": HEAD}),
    ("WO en cours → En transit ; WO clôturé → Arrivé ; WO annulé → alerte.", {"size": 15}),
])

# 9 — Zones de dispatch ---------------------------------------------------------------------
s = content_slide("Chaque zone de Dakar a son dispatch et ses équipages", "Dispatch par zone")
zones = [("Dakar Plateau-Médina", "disp.plateau", ["Hôpital Principal de Dakar", "Clinique de la Madeleine",
                                                  "Centre hospitalier Abass Ndao"]),
         ("Fann-Grand Yoff", "disp.fann", ["CHNU de Fann", "Hôpital d'Enfants Albert Royer",
                                           "Hôpital Général Idrissa Pouye"]),
         ("Pikine-Guédiawaye", "disp.banlieue", ["Hôpital Dalal Jamm"])]
box(s, 3.9, 1.85, 5.5, 0.8, fill=TEAL, name="Région")
text(s, 3.9, 1.85, 5.5, 0.8, "Région de Dakar · Qualification transferts", size=15, bold=True, color=WHITE,
     anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
for i, (z, d, hs) in enumerate(zones):
    x = 0.6 + i * 4.1
    arrow(s, 6.65, 2.65, x + 1.9, 3.05)
    box(s, x, 3.05, 3.8, 3.5, fill=MINT, name=f"Zone {z}")
    text(s, x + 0.3, 3.25, 3.3, 0.45, z, size=17, bold=True, color=TEAL_D)
    text(s, x + 0.3, 3.75, 3.3, 0.7, [f"Dispatch : {d}", "Équipages : 2 ambulanciers"], size=13, color=GREY,
         spacing=2)
    text(s, x + 0.3, 4.6, 3.3, 1.8, hs, size=14, spacing=4)
text(s, 0.6, 6.75, 12, 0.3, "Un groupe couvre sa zone via sys_group_covers_location ; FSM remonte la hiérarchie "
     "des localisations.", size=11, color=GREY)

# 10 — Capacités (graphique) -----------------------------------------------------------------
s = content_slide("Les données de démo comptent 633 lits, dont 95 disponibles", "Données de démonstration")
noms = ["CHNU de Fann", "Hôpital Principal", "Albert Royer", "Idrissa Pouye", "Abass Ndao", "Dalal Jamm",
        "La Madeleine"]
total = [150, 132, 83, 77, 70, 75, 46]
dispo = [20, 23, 9, 12, 14, 8, 9]
cd = CategoryChartData()
cd.categories = noms
cd.add_series("Lits occupés", [t - d for t, d in zip(total, dispo)])
cd.add_series("Lits disponibles", dispo)
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(0.6), Inches(1.8), Inches(8.2), Inches(4.9), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size, ch.legend.font.name = Pt(12), BODY
ch.plots[0].gap_width = 60
for ser, col in zip(ch.series, [LINE, TEAL]):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(col)
ds = ch.series[1]
ds.data_labels.show_value = True
ds.data_labels.position = XL_LABEL_POSITION.INSIDE_BASE
ds.data_labels.font.size, ds.data_labels.font.bold = Pt(11), True
ds.data_labels.font.color.rgb = rgb(WHITE)
ca = ch.category_axis
ca.reverse_order = True
ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(12), BODY
ca.format.line.color.rgb = rgb(LINE)
va = ch.value_axis
va.tick_labels.font.size = Pt(11)
va.tick_labels.font.color.rgb = rgb(GREY)
va.major_gridlines.format.line.color.rgb = rgb("E3ECEA")
va.format.line.fill.background()
text(s, 9.2, 1.9, 3.5, 4.8, [
    ("7 hôpitaux réels", {"bold": True, "size": 18, "color": TEAL_D, "font": HEAD}),
    ("Noms, adresses et standards publics réels.", {"size": 14, "spacing": 12}),
    ("Capacités fictives", {"bold": True, "size": 18, "color": "8A5A12", "font": HEAD}),
    ("Services, lits et équipements sont inventés pour la démonstration.", {"size": 14, "spacing": 12}),
    ("10 spécialités · 6 types d'équipements", {"size": 14, "color": GREY}),
])
s.notes_slide.notes_text_frame.text = "Préciser à l'oral que les capacités sont fictives : aucune donnée réelle de disponibilité."

# 11 — Record producer ------------------------------------------------------------------------
s = content_slide("Un formulaire guidé crée et publie la demande", "Record producer")
secs = [("Origine et urgence", ["Établissement d'origine *", "Service d'origine (filtré)", "Urgence *",
                                "Date souhaitée"]),
        ("Patient", ["Référence pseudonymisée *", "Âge, sexe", "Mobilité *"]),
        ("Besoin et transport", ["Spécialité requise *", "Niveau de soins *", "Mode de transport *",
                                 "Équipements requis"])]
for i, (t, items) in enumerate(secs):
    x = 0.6 + i * 4.1
    box(s, x, 1.95, 3.8, 3.2, fill=MINT, name=f"Section {t}")
    text(s, x + 0.3, 2.15, 3.2, 0.45, t, size=17, bold=True, color=TEAL_D)
    text(s, x + 0.3, 2.75, 3.2, 2.3, items, size=14, spacing=6)
box(s, 0.6, 5.45, 12.1, 1.15, fill=WHITE, line=LINE)
text(s, 0.9, 5.55, 11.6, 1.0, [
    ("☑ Publier immédiatement (cochée par défaut) : le post-insert script publie la demande.", {"size": 15}),
    ("Service Catalog et Service Portal · réservé aux régulateurs · * = obligatoire", {"size": 13, "color": GREY}),
], anchor=MSO_ANCHOR.MIDDLE)

# 12 — Sécurité --------------------------------------------------------------------------------
s = content_slide("Un coordinateur ne voit que son établissement", "Sécurité")
rows = [("Table", "Régulateur", "Coordinateur"),
        ("Référentiel", "Lecture", "Lecture ; modifie son établissement et ses services"),
        ("Demande", "Lecture, création, modification", "Lecture des demandes publiées à son établissement"),
        ("Réponse", "Lecture, retenir une destination", "Lecture et réponse pour son établissement")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.95), Inches(12.1), Inches(2.6)).table
for c, wdt in enumerate([2.4, 4.2, 5.5]):
    tbl.columns[c].width = Inches(wdt)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(TEAL if r == 0 else (MINT if r % 2 else WHITE))
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.runs[0].font.size = Pt(14)
        p.runs[0].font.name = BODY
        p.runs[0].font.bold = r == 0 or c == 0
        p.runs[0].font.color.rgb = rgb(WHITE if r == 0 else INK)
box(s, 0.6, 4.9, 5.9, 1.7, fill="FCF1DF")
text(s, 0.9, 5.05, 5.4, 1.5, [("Action requise", {"bold": True, "size": 16, "color": "8A5A12"}),
                              ("Exécuter le Fix Script « Transfert Patients - créer les ACL » en security_admin.",
                               {"size": 14})])
box(s, 6.8, 4.9, 5.9, 1.7, fill=MINT)
text(s, 7.1, 5.05, 5.4, 1.5, [("Données patient minimales", {"bold": True, "size": 16, "color": TEAL_D}),
                              ("Référence pseudonymisée, âge, sexe, besoin clinique. Jamais de nom ni de "
                               "numéro de sécurité sociale.", {"size": 14})])

# 13 — Feuille de route -------------------------------------------------------------------------
s = content_slide("Priorité : ne plus perdre de lits", "Limites et feuille de route")
cols = [("Maintenant", AMBER, ["Libérer le lit si la demande ou le transport est annulé",
                                "Table « Réservation de lit » pour la traçabilité"]),
        ("Ensuite", TEAL, ["Bloquer le lit dès l'acceptation, avec expiration",
                           "Fusionner transport et remise en une tâche",
                           "Dépendances entre tâches et compétences requises"]),
        ("Plus tard", GREY, ["Portail dédié aux établissements", "Distance dans le choix de destination",
                             "Tableau de bord Performance Analytics"])]
for i, (t, col, items) in enumerate(cols):
    x = 0.6 + i * 4.1
    box(s, x, 1.95, 3.8, 0.7, fill=col, name=f"Horizon {t}")
    text(s, x + 0.3, 1.95, 3.2, 0.7, t, size=18, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x, 2.75, 3.8, 3.85, fill=MINT if i < 2 else "F1F4F4")
    tb = text(s, x + 0.3, 2.95, 3.2, 3.5, items, size=15, spacing=12)

# 14 — Conclusion ----------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
bg(s, TEAL_D)
text(s, 0.8, 1.0, 11.5, 1.0, "Prochaines étapes", size=40, font=HEAD, bold=True, color=WHITE)
nxt = [("Activer la sécurité", "Lancer le Fix Script des ACL en security_admin"),
       ("Tester le parcours", "reg.demo → coord.albertroyer → qualification → disp.fann"),
       ("Valider la feuille de route", "Libération des lits et table de réservation en premier")]
for i, (t, d) in enumerate(nxt):
    y = 2.5 + i * 1.3
    badge(s, 0.8, y, str(i + 1), fill=AMBER, color=TEAL_D)
    text(s, 1.6, y - 0.05, 10.5, 0.45, t, size=20, bold=True, color=WHITE)
    text(s, 1.6, y + 0.45, 10.5, 0.45, d, size=15, color="CFE6E2")
text(s, 0.8, 6.7, 11.5, 0.4, "Documentation complète et scripts de build : dossier TransfertPatients/ du repo.",
     size=12, color="9CC5BF")

prs.save(OUT)
print("écrit :", OUT, "—", len(prs.slides), "diapositives")
